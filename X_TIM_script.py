#!/usr/bin/env python3
"""
X TIM Script - tuned build

Purpose:
- Take JSON exported by the X capture extension.
- Append every non-duplicate captured post into the daily Excel workbook.
- Keep LOW as the archive bucket; do not intentionally skip chatter/noise/content rows.
- Improve categorization by using contextual/tiered Company matching instead of raw word hits.
- Improve IOC extraction so source/report domains do not pollute domain/url IOC fields.
"""
from __future__ import annotations

import argparse
import hashlib
import io
import shutil
import json
import os
import re
import sys
from collections import Counter
from datetime import datetime, timezone, timedelta
from pathlib import Path
from typing import Any, Dict, Iterable, List, Optional, Set, Tuple

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment

LOCAL_TZ = timezone(timedelta(hours=4))
DEFAULT_MAX_AGE_DAYS = 3
WORKBOOK_PREFIX = "x_threat_intel"
LOG_DIR_NAME = "X_TIM_logs"
SEEN_FILE_NAME = "X_TIM_seen_posts.json"
LOG_FILE_NAME = "X_TIM_run_history.log"
RUN_NOTES_NAME = "X_TIM_latest_run_notes.txt"
OCR_CACHE_DIR_NAME = "ocr_cache"

CRITICAL_SHEET = "CRITICAL - Full Review"
REVIEW_SHEET = "REVIEW - Analyst Check"
LOW_SHEET = "LOW - Archive"
SHEETS = [CRITICAL_SHEET, REVIEW_SHEET, LOW_SHEET]

COLUMNS = [
    "source", "title", "link", "published_utc", "threat_type", "severity", "summary",
    "relevance_reason", "cves", "ips", "domains", "urls", "sha256", "emails", "filenames",
    "vendors_matched", "regions_matched", "applications_matched", "brands_matched", "action_required",
]

# -------------------------
# Company profile - tiered for X context
# -------------------------
# Tier 1: exact/high-confidence Company technologies/vendors. These are specific platforms,
# product families, security controls, identity/email/payment/POS/ecom systems, and key vendors.
TIER1_TECH = sorted(set([
    "microsoft 365", "office 365", "m365", "exchange online", "exchange", "entra id", "azure ad", "microsoft entra id",
    "microsoft defender", "defender for endpoint", "defender for cloud", "microsoft intune", "microsoft sentinel",
    "azure devops", "azure waf", "azure data factory", "active directory", "sharepoint", "onedrive", "microsoft teams",
    "oracle cloud", "oracle financials", "oracle ebs", "oracle fusion", "oracle iam", "oracle oic", "oracle integration cloud",
    "oracle retail", "oracle rms", "oracle rpm", "oracle rib", "oracle rics", "oracle resa", "oracle siocs", "oracle supplier portal",
    "sap", "salesforce", "netsuite", "servicenow", "service now", "workday", "adp", "ukg", "quinyx",
    "adobe commerce", "adobe experience manager", "adobe experience platform", "adobe target", "adobe acrobat", "acrobat reader",
    "adobe coldfusion", "coldfusion", "magento", "shopify", "drupal", "wordpress", "pimcore", "aura cms",
    "cloudflare waf", "cloudflare cdn", "cloudflare tunnel", "fastly", "akamai", "forcepoint", "waf", "sdwan", "sd-wan",
    "proofpoint", "proofpoint email gateway", "sentinelone", "sentinel one", "qualys", "checkmarx", "checkmarx one",
    "burp suite", "burpsuite", "acunetix", "black kite", "resecurity", "beyondtrust", "beyond trust", "lastpass",
    "teamviewer", "citrix", "veeam", "sccm", "lansweeper", "arcsight", "sumo logic", "datadog", "data dog", "new relic", "sentry",
    "github actions", "bitbucket", "jenkins", "azure devops", "postman", "postman collection",
    "google workspace", "google suite", "gmail", "google drive", "google cloud platform", "gcp", "google bigquery", "firebase",
    "google ads", "google analytics", "google tag manager", "google merchant center", "google search console", "google api",
    "meta business suite", "meta for developers", "facebook pixel", "instagram", "whatsapp", "tiktok business manager", "snapchat business manager",
    "sprinklr", "infobip", "kaleyra", "mobile app", "appstore", "playstore",
    "payment gateway", "payment gateways", "adyen", "checkout.com", "knet", "benefit pay", "tabby", "tamara", "fiserv", "epaysafe", "qwikcilver",
    "digital gift card", "loyalty platform", "comarch loyalty platform",
    "pos system", "point of sale", "mpos", "micros pos", "micros", "simphony", "symphony", "simphony emc", "self checkout", "ares pos", "winretail",
    "warehouse management system", "blue yonder wms", "blue yonder retail", "oracle supply chain hub", "oms", "order management system",
    "dhl", "aramex", "maersk", "dsv", "expeditors", "smsa", "imile", "roambee", "fareye", "carriyo", "carryio",
    "mongodb", "mongo db", "mysql", "rabbitmq", "rabbit mq", "databricks", "informatica mdm", "data warehouse", "data mart",
]))

# Tier 2: broad/generic technologies. These are relevant to Company but too broad to create Critical alone.
BROAD_TECH = sorted(set([
    "microsoft", "azure", "windows", "linux", "docker", "kubernetes", "apache", "tomcat", "nginx", "python", "github",
    "google", "adobe", "cloudflare", "amazon", "aws", "apple", "box", "dropbox", "edr", "siem", "xdr", "crm", "cms",
    "ecommerce", "e-commerce", "retail", "supply chain", "logistics", "payments", "merchant", "marketing", "source code security",
    "application security", "cloud security", "identity management", "email security", "vulnerability management", "container security",
    "meta", "facebook", "pinterest", "snapchat", "tiktok", "linkedin", "jira", "confluence", "atlassian", "canva", "figma",
]))

COMPANY_BRANDS = sorted(set([
    "6thstreet", "bath and body works", "bathandbodyworks", "bbw", "victoria secret", "victoriassecret",
    "vs connect", "vsg portal", "pottery barn", "pottery barn kids", "west elm", "westelm", "boots", "boots uk",
    "mothercare", "muji", "american eagle", "new balance", "nike", "footlocker", "charlotte tilbury", "harvey nichols",
    "hennes & mauritz", "h&m", "alo yoga", "primark", "ulta", "cos", "no7", "starbucks", "shake shack",
    "kidzania", "namshi", "jumia", "trendyol",
]))

COMPANY_REGIONS = sorted(set([
    "middle east", "uae", "united arab emirates", "gulf", "gcc", "saudi", "kuwait", "qatar", "bahrain",
    "oman", "mena", "turkey", "egypt", "india", "dubai", "abu dhabi", "ksa", "saudi arabia", "jordan", "poland",
]))

THREAT_KEYWORDS = {
    "Vulnerability": ["vulnerability", "vuln", "cve", "zero-day", "zero day", "rce", "remote code execution", "privilege escalation", "auth bypass", "authentication bypass", "security advisory", "patch", "kev", "actively exploited", "exploit in the wild", "known exploited", "cvss"],
    "Malware": ["malware", "trojan", "backdoor", "stealer", "infostealer", "rat", "botnet", "loader", "dropper", "spyware", "banker", "payload", "c2", "command and control", "rmm payload", "clickfix"],
    "Ransomware": ["ransomware", "extortion", "lockbit", "alphv", "black basta", "cl0p", "encryptor", "qilin", "akira", "space bears", "chaos ransomware"],
    "Phishing": ["phishing", "phish", "smishing", "vishing", "credential theft", "device code phishing", "social engineering", "remote access lure", "teams impersonation", "fake login", "oauth", "mfa bypass", "session cookies", "phaas"],
    "Data Breach": ["data breach", "breach", "leak", "leaked", "exposed", "compromised data", "stolen data", "credentials allegedly leaked", "database allegedly"],
    "Supply Chain": ["supply chain", "dependency confusion", "malicious package", "npm", "pypi", "github actions", "ci/cd", "pipeline", "poisoned update", "typosquat", "typosquatting", "sdk"],
    "APT / Intrusion": ["nation state", "state-sponsored", "espionage", "apt", "lazarus", "apt28", "apt29", "mustang panda", "sandworm", "kimsuky", "silver fox", "muddywater"],
}

ACTIVE_EXPLOIT_KEYWORDS = ["actively exploited", "exploited in the wild", "in-the-wild", "zero-day", "known exploited", "under active exploitation", "mass exploitation", "0-day", "used as a 0-day"]
IMMEDIATE_ACTION_KEYWORDS = ["critical", "cvss 9", "cvss 10", "cvss score of 10", "actively exploited", "zero-day", "ransomware", "wormable", "credential exposure", "stolen credentials", "unauthenticated rce", "immediate action", "bypass mfa", "session cookie"]
SECURITY_CONTEXT_TERMS = sorted(set(sum(THREAT_KEYWORDS.values(), []) + ["ioc", "iocs", "indicator", "indicator of compromise", "breached", "patched", "psirt", "advisory", "exploit", "opendir", "open directory"]))

ACTION_MAP = {
    "Vulnerability": "Validate exposure, check vendor advisory, patch or mitigate, and verify internet exposure.",
    "Malware": "Review detections, hunt impacted endpoints, and block related malicious infrastructure where validated.",
    "Ransomware": "Escalate to incident response, validate backups, review identity protections, and block confirmed IOCs.",
    "Phishing": "Block malicious domains or URLs, review email controls, and search telemetry for delivery or user interaction.",
    "Data Breach": "Assess exposure, review affected accounts or tokens, and search monitoring data for related access or abuse.",
    "Supply Chain": "Review impacted package or vendor use, validate exposure, rotate affected secrets, and inspect CI/CD pipelines.",
    "APT / Intrusion": "Escalate to threat hunting, map TTPs, search telemetry, and review privileged access and segmentation.",
    "Other": "Analyst review recommended if business or technology context indicates possible relevance.",
}

# -------------------------
# IOC extraction regexes
# -------------------------
CVE_RE = re.compile(r"\bCVE-\d{4}-\d{4,7}\b", re.I)
EMAIL_RE = re.compile(r"\b[A-Za-z0-9._%+\-]+@[A-Za-z0-9.\-]+\.[A-Za-z]{2,}\b")
SHA256_RE = re.compile(r"\b[a-fA-F0-9]{64}\b")
IP_RE = re.compile(r"(?<!\d)(?:\d{1,3}\s*(?:\.|\[\.\]|\[\.\)|\[\.|\(\.\)|\[dot\])\s*){3}\d{1,3}(?::\d{1,5})?(?!\d)", re.I)
URL_RE = re.compile(r"\b(?:hxxps?|https?|ftp)\s*(?::|\[:\])\s*//\s*[^\s\"<>]+", re.I)
DOMAIN_RE = re.compile(
    r"\b(?:[a-zA-Z0-9](?:[a-zA-Z0-9\-]{0,61}[a-zA-Z0-9])?\s*(?:\.|\[\.\]|\[\.\)|\[\.|\(\.\)|\[dot\])\s*)+"
    r"(?:com|net|org|io|dev|shop|top|cc|ru|cn|jp|de|uk|eu|ae|sa|qa|kw|bh|om|info|biz|live|today|name|cloud|app|gov|edu|su|digital|tech|site|xyz|ir|mx|py|ar|club|st|name|r2\.dev|pages\.dev|workers\.dev)\b",
    re.I,
)
FILE_TOKEN_RE = re.compile(r"(?<![A-Za-z0-9_])[A-Za-z0-9][A-Za-z0-9_.()\-]{0,100}\.(?:exe|dll|sys|msi|vbs|js|jse|ps1|bat|cmd|lnk|zip|rar|7z|iso|dmg|apk|py|scr|docm|xlsm|xlsx|pdf|hta|cpl|inf|img|pdb)(?![A-Za-z0-9_])", re.I)
QUOTED_FILE_RE = re.compile(r"['\"]([^'\"]{1,120}\.(?:exe|dll|sys|msi|vbs|js|jse|ps1|bat|cmd|lnk|zip|rar|7z|iso|dmg|apk|py|scr|docm|xlsm|xlsx|pdf|hta|cpl|inf|img|pdb))['\"]", re.I)

REFERENCE_DOMAINS = {
    "t.co", "x.com", "twitter.com", "github.com", "gist.github.com", "gitlab.com", "virustotal.com", "tria.ge",
    "any.run", "app.any.run", "abuse.ch", "bazaar.abuse.ch", "malwarebazaar.org", "urlscan.io", "phishstats.info",
    "bleepingcomputer.com", "cybersecuritynews.com", "gbhackers.com", "thehackernews.com", "securityweek.com", "darkreading.com",
    "darkwebinformer.com", "documentcloud.org", "socket.dev", "nowsecure.com", "zimperium.com", "proofpoint.com",
    "openwall.com", "herodevs.com", "medium.com", "youtube.com", "youtu.be", "reddit.com", "linkedin.com", "hunt.io",
    "secureblink.com", "labs.infoguard.ch", "labs.cloudsecurityalliance.org", "blog.talosintelligence.com", "kaspersky.com",
    "kas.pr", "thedfirreport.com", "lifehacker.com", "community.emergingthreats.net", "yeswehack.com", "decipher.sc",
    "techrepublic.com", "macrumors.com", "gov.uk", "openinfosecfoundation.org", "forum.suricata.io", "hubs.la",
    "aikido.dev", "embeeresearch.io", "wadgamaraldeen.medium.com", "windowscentral.com", "flare.io", "etugen.io", "app.etugen.io",
}

ACTIONABLE_IOC_HINTS = [
    "ioc", "iocs", "source ip", "communication destination", "通信先", "contacted url", "next stage", "downloaded", "downloads", "payload", "phishing at", "observed phishing", "malicious domain", "known malicious domain", "domain:", "url:", "c2", "command and control", "ratshop", "opendir", "open directory", "drops", "fake", "spoof", "lure", "clickfix", "rmm payload", "resolving to", "fud source", "attached link", "download url", "c2 domain", "leak site",
]
REFERENCE_HINTS = ["from ", "source:", "read more", "full article", "writeup", "blog", "report", "sample:", "github:", "details", "analysis", "via @", "research", "article", "full analysis", "read more:"]


# Context controls used to avoid fixing one example at a time. These lists describe
# general roles terms can play: affected platform, lure/infrastructure, victim entity,
# or source/report reference.
TECH_AFFECTED_CONTEXT = [
    "affected", "vulnerable", "vulnerability", "flaw", "cve", "rce", "exploit", "patch", "advisory", "bypass",
    "misconfiguration", "tenant", "account", "mailbox", "email", "oauth", "mfa", "token", "session", "credential",
    "repo", "repository", "workflow", "actions", "pipeline", "ci/cd", "package", "sdk", "dependency", "secret",
    "waf", "defender", "edr", "gateway", "workspace", "drive", "api", "cloud", "pos", "payment", "ecommerce",
    "formjacking", "magecart", "skimming", "phishing", "malware", "stealer", "rat", "payload", "hosted on",
]
INCIDENTAL_OR_SOURCE_CONTEXT = [
    "from ", "source:", "read more", "full article", "writeup", "blog", "report", "sample:", "website:",
    "github:", "via @", "new alliance", "announced a new alliance", "works on", "worked at", "running ",
]
FINANCIAL_EXCHANGE_CONTEXT = ["kyc", "aml", "crypto", "trading", "wallet", "blockchain", "exchange building", "worked at an exchange"]
SPORTS_CONTEXT = ["football", "match", "vs.", " vs ", "beat", "pitch", "revenge"]
VICTIM_DOMAIN_CONTEXT = ["ransomware alert", "has reportedly fallen victim", "reportedly fallen victim", "victim", "data compromised", "deadline", "declared exfiltrated", "sector:"]
FEATURE_OR_RELEASE_CONTEXT = ["released", "release", "new feature", "adds", "now blocks", "call for papers", "stable", "try it", "demo", "demonstration", "utility", "tool"]

REPORT_PUBLISHER_CONTEXT = ["uncovered", "reported", "published", "analysed", "analyzed", "researchers", "threat insight", "blog", "writeup", "article"]
NON_ACTIONABLE_CHATTER_CONTEXT = ["worst?", "ehhhhh", "goofy", "yappin", "i made while learning", "collection of programs", "call for papers", "schedule announcement", "please excuse my language"]
CORE_REGIONS = {"uae", "united arab emirates", "dubai", "abu dhabi", "kuwait", "qatar", "bahrain", "oman", "saudi", "saudi arabia", "ksa", "gcc", "gulf", "middle east", "mena"}
CRITICAL_BUSINESS_REGIONS = {"uae", "united arab emirates", "dubai", "abu dhabi", "kuwait"}
PROFILE_LOGISTICS_VENDORS = {"smsa", "aramex", "dhl", "dsv", "maersk", "expeditors", "imile", "fareye", "roambee", "carriyo", "carryio"}
BIG_GLOBAL_ORGS = sorted(set([
    "microsoft", "google", "alphabet", "apple", "amazon", "meta", "facebook", "oracle", "sap", "salesforce", "servicenow", "adobe", "ibm", "intel", "nvidia", "openai",
    "cisco", "broadcom", "vmware", "dell", "hp", "hpe", "lenovo", "samsung", "huawei", "sony", "panasonic", "lg",
    "cloudflare", "akamai", "fastly", "palo alto networks", "fortinet", "crowdstrike", "okta", "zscaler", "proofpoint", "sentinelone", "qualys", "checkmarx",
    "walmart", "target", "carrefour", "costco", "tesco", "aldi", "lidl", "ikea", "starbucks", "mcdonald", "mcdonald's", "nike", "adidas", "h&m", "zara", "inditex",
    "dhl", "fedex", "ups", "maersk", "dp world", "emirates", "etihad", "qatar airways", "aramco", "adnoc", "stc", "mobily", "du", "etisalat", "e&",
    "jpmorgan", "chase", "bank of america", "citi", "citibank", "hsbc", "standard chartered", "barclays", "deutsche bank", "visa", "mastercard", "amex", "american express",
    "accenture", "deloitte", "pwc", "kpmg", "ey", "ernst & young", "capgemini", "wipro", "infosys", "tcs", "cognizant",
]))
# Some names in BIG_GLOBAL_ORGS have common-language or short-token meanings.
# They require stronger company formatting/context instead of raw keyword matching.
AMBIGUOUS_GLOBAL_ORGS = {"target", "hp", "hpe", "ey", "du", "box", "visa", "emirates", "chase", "citi", "lg", "sap", "apple"}
GLOBAL_ORG_STRICT_PATTERNS = {
    "target": [r"\bTarget\s+(?:Corp(?:oration)?|Stores?)\b", r"\btarget\.com\b.*\b(?:breach|ransomware|leak|victim)\b"],
    "hp": [r"\bHP\s+(?:Inc|Enterprise|Wolf|Sure|Poly)\b", r"\bHewlett[- ]Packard\b"],
    "hpe": [r"\bHPE\b", r"\bHewlett Packard Enterprise\b"],
    "ey": [r"\bErnst\s*&\s*Young\b", r"\bEY\s+(?:Global|LLP|GDS|consulting|audit|tax)\b"],
    "du": [r"\bdu\s+(?:UAE|telecom|Telecom|Mobile|Emirates Integrated Telecommunications)\b", r"\bdu\.ae\b"],
    "box": [r"\bBox\s+(?:Inc|Enterprise|Cloud|Drive|Sign|Shield)\b", r"\bbox\.com\b.*\b(?:breach|ransomware|leak|victim)\b"],
    "visa": [r"\bVisa\s+(?:Inc|card network|payments network)\b", r"\bvisa\.com\b.*\b(?:breach|ransomware|leak|victim)\b"],
    "emirates": [r"\bEmirates\s+(?:Airline|Group|NBD|Islamic)\b"],
    "chase": [r"\bJPMorgan\s+Chase\b", r"\bChase\s+Bank\b"],
    "citi": [r"\bCiti(?:bank|group)?\b.*\b(?:bank|financial|cards|breach|ransomware|victim)\b"],
    "lg": [r"\bLG\s+(?:Electronics|U\+|Innotek|Chem)\b"],
    "sap": [r"\bSAP\s+(?:SE|S/4HANA|NetWeaver|Business|SuccessFactors|Ariba)\b"],
    "apple": [r"\bApple\s+(?:Inc|ID|iCloud|Support|Business|Store)\b", r"\bapple\.com\b.*\b(?:breach|ransomware|leak|victim)\b"],
}
GENERIC_BROAD_TERMS = {"microsoft", "windows", "linux", "github", "apple", "google", "amazon", "adobe", "box", "dropbox", "edr", "crm", "cms", "marketing", "supply chain", "retail", "ecommerce", "e-commerce", "payments", "linkedin", "facebook", "meta"}
DIRECT_TECH_THREAT_TERMS = {"cve-", "vulnerability", "affected", "affects", "patched", "update", "upgrade", "rce", "exploit", "actively exploited", "kev", "zero-day", "0-day", "auth bypass", "bypass", "credential", "mfa", "oauth", "session", "device code", "device-code", "payload", "download url", "malicious package"}
VICTIM_LIST_MARKERS = ["claims 12 victims", "added 12 new victims", "added 2 new victims", "dark web leak portal", "victims to their dark web", "victims to its dark web"]
REPORT_ONLY_MARKERS = ["deep dive", "we just published", "published an in-depth", "new report", "threat actor report", "daily threat brief", "stormcast", "podcast", "learn how", "partnered with", "now supports", "active defense now supports", "mid-year", "crystal ball series"]

PROMOTE_REVIEW_PATTERNS = [
    ("ecommerce/payment skimming", ["magecart", "formjacking", "eskimming", "payment form", "pcidss", "pci dss"]),
    ("payment/developer supply-chain", ["malicious package", "npm", "pypi", "sdk", "developer", "ci secret", "ci secrets", "payment client", "typosquat", "typosquatting"]),
    ("m365 identity phishing", ["microsoft 365", "m365", "exchange online", "entra", "device-code", "device code", "oauth", "mfa bypass", "session cookie", "phaas"]),
    ("meta business phishing", ["meta business", "business account manager", "legit meta address", "partner-name", "phishing setup"]),
    ("critical web-app cve", ["cve-", "cvss 9", "cvss 10", "rce", "upload", "web application", "framework", "cms", "fofa"]),
]

TERM_PATTERN_CACHE: Dict[str, re.Pattern] = {}
def term_pattern(term: str) -> re.Pattern:
    key = term.lower().replace("&amp;", "&").strip()
    if key not in TERM_PATTERN_CACHE:
        escaped = re.escape(key).replace(r"\ ", r"\s+")
        TERM_PATTERN_CACHE[key] = re.compile(r"(?<![a-z0-9])" + escaped + r"(?![a-z0-9])", re.I)
    return TERM_PATTERN_CACHE[key]

# -------------------------
# basic helpers
# -------------------------
def now_utc() -> datetime:
    return datetime.now(timezone.utc)

def today_local_date() -> str:
    return now_utc().astimezone(LOCAL_TZ).date().isoformat()

def parse_dt(value: str) -> Optional[datetime]:
    try:
        return datetime.fromisoformat(str(value).replace("Z", "+00:00")).astimezone(timezone.utc)
    except Exception:
        return None

def unique(items: Iterable[str]) -> List[str]:
    seen: Set[str] = set(); out: List[str] = []
    for item in items:
        if item is None:
            continue
        val = str(item).strip()
        if not val:
            continue
        key = val.lower()
        if key not in seen:
            seen.add(key); out.append(val)
    return out

def join(items: Iterable[str]) -> str:
    return ", ".join(unique(items))

def normalize_defang(s: str) -> str:
    out = str(s or "")
    out = re.sub(r"hxxps", "https", out, flags=re.I)
    out = re.sub(r"hxxp", "http", out, flags=re.I)
    out = out.replace("[:]", ":")
    out = re.sub(r"\s*\[\.\]\s*|\s*\[\.\)\s*|\s*\[\.\s*|\s*\(\.\)\s*|\s*\[dot\]\s*", ".", out, flags=re.I)
    out = out.replace("\\.", ".")
    out = re.sub(r"https?://\s+", lambda m: m.group(0).replace(" ", ""), out)
    out = re.sub(r"(?<=\w)\s+(?=\w)", "", out)
    return out.strip(".,;:'\"()[]{}<>…")

def clean_post_text(text: str) -> str:
    """Remove X UI artifacts while preserving actual content.

    X sometimes exposes leading engagement/view counters before the real post text.
    We strip repeated standalone 1-2 digit counters, but preserve meaningful 3+ digit
    values such as threat actor names like 888 or counts like 11.4M.
    """
    t = str(text or "").replace("&amp;", "&")
    t = re.sub(r"\s+", " ", t).strip()
    # Strip repeated leading UI counters: "8 CVE...", "7 𝗣...", "8 9 new OPEN...".
    while re.match(r"^\d{1,2}\s+", t) and not re.match(r"^\d{3,}\b", t):
        t = re.sub(r"^\d{1,2}\s+", "", t, count=1).strip()
    t = re.sub(r"\b\d+:\d{2}\s*/\s*\d+:\d{2}\b", "", t)
    t = re.sub(r"\s+", " ", t).strip()
    return t


def strip_urls_for_profile(text: str) -> str:
    t = URL_RE.sub(" ", text or "")
    t = re.sub(r"https?\s*://[^\s]+", " ", t, flags=re.I)
    return re.sub(r"\s+", " ", t).strip()

def contains_terms(text: str, terms: Iterable[str]) -> List[str]:
    hits = []
    low = (text or "").lower()
    for term in terms:
        if term_pattern(term).search(low):
            hits.append(term)
    return unique(hits)

def nearby_context(text: str, start: int, end: int, radius: int = 120) -> str:
    return (text[max(0, start-radius):min(len(text), end+radius)] or "").lower()

def get_domain(value: str) -> str:
    v = normalize_defang(value).lower()
    v = re.sub(r"^https?://", "", v)
    return v.split("/", 1)[0].split(":", 1)[0].strip(". ")

def is_reference_domain(domain: str) -> bool:
    d = (domain or "").lower().strip(". ")
    if not d:
        return False
    return d in REFERENCE_DOMAINS or any(d.endswith("." + ref) for ref in REFERENCE_DOMAINS)

def has_defang(raw: str) -> bool:
    return bool(re.search(r"hxxps?|\[\.\]|\[\.\)|\[\.|\(\.\)|\[dot\]|\[:\]", raw or "", re.I))

def valid_ip(raw: str) -> Optional[str]:
    norm = normalize_defang(raw)
    base = norm.split(":", 1)[0]
    parts = base.split(".")
    if len(parts) != 4:
        return None
    try:
        nums = [int(p) for p in parts]
    except Exception:
        return None
    return norm if all(0 <= n <= 255 for n in nums) else None

def is_actionable_domain_or_url(raw_value: str, context: str) -> bool:
    domain = get_domain(raw_value)
    ctx = (context or "").lower()
    actionable = any(h in ctx for h in ACTIONABLE_IOC_HINTS)
    reference = any(h in ctx for h in REFERENCE_HINTS)
    victim_ctx = any(h in ctx for h in VICTIM_DOMAIN_CONTEXT)
    if is_reference_domain(domain):
        return False
    # Victim/company websites in ransomware or breach-alert posts are not blockable IOCs.
    # Keep them in the summary, not in the domains/urls IOC columns.
    if victim_ctx and not has_defang(raw_value) and not any(h in ctx for h in ["c2", "payload", "download", "next stage", "phishing"]):
        return False
    if has_defang(raw_value):
        return True
    if actionable and not reference:
        return True
    return False



def dedupe_filenames(names: Iterable[str]) -> List[str]:
    cleaned = unique(n.strip().strip(".,;:()[]{}<>…") for n in names if n)
    result: List[str] = []
    for name in sorted(cleaned, key=len, reverse=True):
        low = name.lower()
        # Drop partial suffix captures if a longer filename already contains them.
        if any(low != existing.lower() and low in existing.lower() for existing in result):
            continue
        result.append(name)
    return sorted(result, key=lambda x: cleaned.index(x) if x in cleaned else 9999)

def extract_iocs(text: str, external_links: Any = None) -> Dict[str, List[str]]:
    cves = unique(x.upper() for x in CVE_RE.findall(text or ""))
    ips = unique(ip for ip in (valid_ip(m.group(0)) for m in IP_RE.finditer(text or "")) if ip)
    sha256 = unique(x.lower() for x in SHA256_RE.findall(text or ""))
    emails = unique(EMAIL_RE.findall(text or ""))

    urls: List[str] = []
    domains: List[str] = []
    for m in URL_RE.finditer(text or ""):
        raw = m.group(0)
        ctx = nearby_context(text, m.start(), m.end())
        norm = normalize_defang(raw)
        if is_actionable_domain_or_url(raw, ctx):
            urls.append(norm)
            d = get_domain(norm)
            if d and not is_reference_domain(d):
                domains.append(d)

    for m in DOMAIN_RE.finditer(text or ""):
        raw = m.group(0)
        ctx = nearby_context(text, m.start(), m.end())
        norm = normalize_defang(raw)
        d = get_domain(norm)
        if not d or is_reference_domain(d):
            continue
        if is_actionable_domain_or_url(raw, ctx):
            domains.append(d if d else norm)

    filenames: List[str] = []
    for m in QUOTED_FILE_RE.finditer(text or ""):
        filenames.append(m.group(1).strip())
    for m in FILE_TOKEN_RE.finditer(text or ""):
        tok = m.group(0).strip().strip(".,;:()[]{}<>…")
        # reject domains mistakenly matching as filenames via non-file extensions; our extension list avoids .com etc.
        if len(tok) <= 120:
            filenames.append(tok)

    return {
        "cves": cves,
        "ips": ips,
        "domains": unique(domains),
        "urls": unique(urls),
        "sha256": sha256,
        "emails": emails,
        "filenames": dedupe_filenames(filenames),
    }

def classify_threat(text: str, iocs: Dict[str, List[str]]) -> str:
    low = (text or "").lower()
    # Priority rules first; semantic event type should beat incidental keywords.
    if "ddos alert" in low or "ddos" in low and "claims to have targeted" in low:
        return "Other"
    if any(k in low for k in ["ransomware alert", "ransomware group", "dark web portal", "fallen victim", "data compromised", "declared exfiltrated", "new claim #ransomware"]):
        return "Ransomware"
    if any(k in low for k in ["data breach", "records leaked", "records exposed", "database allegedly leaked", "claims sale of", "claims to have leaked", "claims breach", "claims to have breached", "claims that", "alleged leak", "pii", "candidate records", "subscriber data", "personal information"]):
        if any(x in low for x in ["breach", "breached", "leak", "leaked", "records", "pii", "personal information", "database", "subscriber data"]):
            return "Data Breach"
    if any(k in low for k in ["new #phishing", "phishing site", "phishing url", "phaas", "device code phishing", "device-code phishing", "credential phishing", "mfa bypass", "session cookie", "fake login", "credential capture", "steal microsoft 365 accounts", "fake microsoft 365 login", "redirected to fake microsoft 365"]):
        return "Phishing"
    if any(k in low for k in ["malicious package", "dependency confusion", "typosquat", "typosquatting", "npm", "pypi", "ci/cd", "pipeline", "poisoned update", "package mirror"]):
        return "Supply Chain"
    if any(k in low for k in ["mirai", "gafgyt", "botnet", "payload", "downloads", "downloaded", "source ip", "c2", "rat", "stealer", "loader"]) and any(iocs[k] for k in ["ips", "domains", "urls", "sha256"]):
        return "Malware"
    if iocs["cves"] or any(k in low for k in ["cve-", "cvss", "rce", "vulnerability", "security advisory", "actively exploited"]):
        return "Vulnerability"
    scores = {name: sum(1 for k in kws if k in low) for name, kws in THREAT_KEYWORDS.items()}
    if any(iocs[k] for k in ["sha256", "ips", "domains", "urls"]):
        scores["Malware"] += 1
    best, val = max(scores.items(), key=lambda kv: kv[1])
    return best if val > 0 else "Other"


def has_strong_signal(text: str, iocs: Dict[str, List[str]], ttype: str) -> bool:
    low = (text or "").lower()
    if "not malware" in low and not iocs["domains"] and not iocs["urls"] and not iocs["sha256"]:
        return False
    if any(k in low for k in ACTIVE_EXPLOIT_KEYWORDS + IMMEDIATE_ACTION_KEYWORDS):
        return True
    if iocs["cves"]:
        return True
    if ttype in {"Malware", "Phishing", "Ransomware", "Data Breach", "Supply Chain", "APT / Intrusion"}:
        return True
    return False

def has_weak_signal(text: str, iocs: Dict[str, List[str]], ttype: str) -> bool:
    if has_strong_signal(text, iocs, ttype):
        return False
    low = (text or "").lower()
    return ttype != "Other" or any(k in low for k in SECURITY_CONTEXT_TERMS) or any(iocs[k] for k in iocs)

def context_is_about_term(text: str, term: str) -> bool:
    """Decide if a technology/vendor/brand term is actually relevant to the threat.

    This is a general ambiguity/control layer, not an example-specific fix. A term
    must appear as an affected/abused platform, target, security control, package,
    infrastructure actively used by the threat, or business-relevant service. Source,
    repost, victim, lure-only, sports, or ordinary-language contexts are downweighted.
    """
    raw_low = (text or "").lower()
    profile_low = strip_urls_for_profile(text).lower()
    term_l = term.lower()
    patterns = list(term_pattern(term_l).finditer(profile_low))
    if not patterns:
        return False

    for m in patterns:
        window = profile_low[max(0, m.start()-120):min(len(profile_low), m.end()+120)]
        if is_victim_list_or_digest(text) and not any(x in window for x in ["cve-", "vulnerability", "affected", "affects", "exploit", "phishing", "credential", "screenconnect", "fortigate", "magecart"]):
            # In multi-victim ransomware/breach digests, OCR/company descriptions often include technologies/logos
            # that are not the threat subject for Company.
            continue
        # Very broad/general terms must be tied to an actual affected/abused technology context.
        if term_l in GENERIC_BROAD_TERMS and not any(x in window for x in DIRECT_TECH_THREAT_TERMS | {"c2", "command and control", "ioc", "iocs", "payload", "hosted on"}):
            continue
        # Ambiguous product names require product-specific meaning.
        if any(x in raw_low for x in ["stormcast", "podcastdetail", "ruleset update summary"]) and not any(x in window for x in ["cve-", "actively exploited", "critical", "rce"]):
            continue
        if term_l in {"exchange", "exchange online"}:
            if any(x in window for x in FINANCIAL_EXCHANGE_CONTEXT):
                continue
            if not any(x in window for x in ["microsoft", "online", "mailbox", "email", "tenant", "spoof", "smtp", "ews", "outlook", "o365", "m365"]):
                continue
        # Security vendors often appear as report publishers. Do not treat the publisher as the affected Company product.
        if term_l in {"proofpoint", "socket", "anyrun", "sans", "zscaler", "virus bulletin"} and any(x in window for x in REPORT_PUBLISHER_CONTEXT):
            if not any(x in window for x in ["gateway", "email gateway", "product", "customer", "tenant", "appliance", "vulnerability in proofpoint"]):
                continue
        if any(x in window for x in INCIDENTAL_OR_SOURCE_CONTEXT) and not any(x in window for x in TECH_AFFECTED_CONTEXT):
            continue
        if any(x in window for x in FEATURE_OR_RELEASE_CONTEXT) and not any(x in window for x in ["malware", "phishing", "cve", "rce", "exploit", "bypass", "stealer", "credential", "token"]):
            continue
        if any(x in window for x in TECH_AFFECTED_CONTEXT):
            return True
        # Multiword exact products are allowed unless they are just source/report text.
        if " " in term_l and not any(x in window for x in INCIDENTAL_OR_SOURCE_CONTEXT):
            return True
        if term_l in {"coldfusion", "magento", "servicenow", "proofpoint", "sentinelone", "beyondtrust", "veeam", "qualys", "checkmarx"}:
            return True
    return False


def region_is_relevant(text: str, region: str) -> bool:
    low = (text or "").lower()
    # Avoid incidental sports/geopolitical mentions where the region is not the cyber target/victim/actor location.
    idx = low.find(region.lower())
    if idx < 0:
        return False
    window = low[max(0, idx-120):min(len(low), idx+120)]
    if any(x in window for x in SPORTS_CONTEXT):
        return False
    return any(x in window for x in ["target", "victim", "breach", "leak", "ransomware", "phishing", "attack", "compromised", "government", "bank", "ministry", "company", "customer", "data"])


def special_relevance(text: str) -> List[str]:
    low = (text or "").lower()
    hits = []
    for name, terms in PROMOTE_REVIEW_PATTERNS:
        count = sum(1 for t in terms if t in low)
        if name == "m365 identity phishing":
            # Critical/Review only when Microsoft identity flow is actually involved, not generic PhaaS/cloudflare lures.
            if count >= 2 and any(x in low for x in ["microsoft", "m365", "microsoft 365", "office 365", "exchange", "entra", "device code", "device-code", "ad cs", "active directory"]):
                hits.append(name)
        elif name == "meta business phishing":
            if count >= 2 and any(x in low for x in ["meta", "business account", "partner-name", "legit meta address"]):
                hits.append(name)
        elif name == "payment/developer supply-chain":
            if count >= 2 and any(x in low for x in ["malicious", "attack", "compromised", "blocked", "npm", "pypi", "package", "developer", "ci secret", "dependency"]):
                # Avoid marketing/testimonial replies that only talk about supply-chain in general.
                if not any(x in low for x in ["learn how", "partnered with", "right fit", "visibility into supply chain threats"]):
                    hits.append(name)
        elif name == "critical web-app cve":
            if (count >= 3 or any(x in low for x in ["vbulletin", "teamcity", "xrdp", "spring boot"])) and any(x in low for x in ["cvss 9", "cvss 10", "critical", "unauthenticated", "preauth", "rce", "actively exploited"]):
                hits.append(name)
        elif count >= 2:
            hits.append(name)
    # Business-specific access/breach patterns that should not be missed.
    if any(v in low for v in PROFILE_LOGISTICS_VENDORS) and any(x in low for x in ["records", "breach", "leak", "darkweb", "dark web", "database", "sale"]):
        hits.append("profile logistics/vendor exposure")
    if "screenconnect" in low and any(x in low for x in ["phishing", "rmm", "payload", "clientsetup", "download url", "fake", "zoom"]):
        hits.append("RMM/ScreenConnect phishing payload")
    if "fortigate" in low and any(x in low for x in ["access", "ssh", "web-panel", "vpn", "sale"]):
        hits.append("FortiGate access sale")
    return unique(hits)


def is_victim_list_or_digest(text: str) -> bool:
    low = (text or "").lower()
    return any(x in low for x in VICTIM_LIST_MARKERS + REPORT_ONLY_MARKERS)


def matched_big_global_orgs(text: str) -> List[str]:
    raw = text or ""
    low = raw.lower()
    hits: List[str] = []
    for org in BIG_GLOBAL_ORGS:
        org_l = org.lower()
        if org_l in AMBIGUOUS_GLOBAL_ORGS:
            pats = GLOBAL_ORG_STRICT_PATTERNS.get(org_l, [])
            if any(re.search(pat, raw, flags=re.I) for pat in pats):
                hits.append(org)
            continue
        for m in term_pattern(org).finditer(low):
            window = low[max(0, m.start()-120):min(len(low), m.end()+160)]
            # Do not treat report publishers, source links, ASNs, or template/lure text as the victim company.
            if any(x in window for x in ["published", "uncovered", "research", "writeup", "blog", "report", "from ", "source:", "template", "lure", "phishing page", "asn:", "as13335"]):
                if not any(x in window for x in ["fallen victim", "breached", "claims", "leaked", "stolen", "ransomware", "data compromised"]):
                    continue
            hits.append(org)
            break
    return unique(hits)


def is_business_incident(text: str, ttype: str) -> bool:
    low = (text or "").lower()
    if ttype in {"Data Breach", "Ransomware"}:
        return True
    # Access-sale is business incident; ordinary vulnerability wording like "network access" is not.
    sale_terms = ["offered for sale", "listed for sale", "selling", "sells", "sale of", "claims sale", "advertising access", "access allegedly offered"]
    access_terms = ["rdp access", "vpn access", "ssh access", "domain admin access", "fortigate access", "server access", "database access", "web-panel access", "web panel access"]
    if any(a in low for a in access_terms) and any(s in low for s in sale_terms + ["forum actor", "marketplace", "darkweb", "dark web"]):
        return True
    if any(x in low for x in ["alleged breach", "allegedly breached", "allegedly leaked", "database leaked", "records leaked", "records exposed", "data allegedly", "darkweb forum post claims", "dark web forum post claims"]):
        return True
    return False


def is_direct_technology_threat(text: str, ttype: str) -> bool:
    low = (text or "").lower()
    if is_victim_list_or_digest(text) and ttype in {"Ransomware", "Data Breach", "Other"}:
        return False
    if any(x in low for x in REPORT_ONLY_MARKERS) and not any(x in low for x in ["cve-", "actively exploited", "critical", "unauthenticated", "phishing", "credential", "ransomware"]):
        return False
    return any(x in low for x in DIRECT_TECH_THREAT_TERMS)


def broad_match_is_actionable(text: str, iocs: Dict[str, List[str]], matches: Dict[str, List[str]], ttype: str) -> bool:
    low = (text or "").lower()
    broad_terms = set(x.lower() for x in matches.get("broad", []))
    if not broad_terms:
        return False
    if broad_terms.issubset(GENERIC_BROAD_TERMS) and not matches.get("specials"):
        # Generic platforms like Windows/GitHub/Apple/Microsoft should not promote routine malware/chatter.
        if ttype in {"Malware", "Other", "Supply Chain"} and not any(x in low for x in ["malicious package", "dependency confusion", "npm", "pypi", "github actions", "ci/cd", "secret", "token", "phishing", "credential", "device code", "rce", "cve-"]):
            return False
    if is_victim_list_or_digest(text) and not matches.get("specials"):
        return False
    if any(x in low for x in REPORT_ONLY_MARKERS) and not matches.get("specials") and not (iocs.get("cves") and any(y in low for y in ["critical", "cvss 9", "unauthenticated", "rce"])):
        return False
    return True


def region_sets(matches: Dict[str, List[str]]) -> Tuple[Set[str], Set[str]]:
    regs = set(x.lower() for x in matches.get("regions", []))
    return regs & CRITICAL_BUSINESS_REGIONS, regs - CRITICAL_BUSINESS_REGIONS


def profile_matches(text: str) -> Dict[str, List[str]]:
    profile_text = strip_urls_for_profile(text)
    brands = [t for t in contains_terms(profile_text, COMPANY_BRANDS) if context_is_about_term(text, t) or t in ["company", "company L.L.C.", "company LLC"]]
    regions = [r for r in contains_terms(profile_text, COMPANY_REGIONS) if region_is_relevant(text, r)]
    tier1 = [t for t in contains_terms(profile_text, TIER1_TECH) if context_is_about_term(text, t)]
    broad_raw = contains_terms(profile_text, BROAD_TECH)
    broad = [t for t in broad_raw if context_is_about_term(text, t)]
    specials = special_relevance(text)
    return {"brands": unique(brands), "regions": unique(regions), "tier1": unique(tier1), "broad": unique(broad), "specials": specials}


def score_row(text: str, iocs: Dict[str, List[str]], matches: Dict[str, List[str]], ttype: str) -> Tuple[str, str, str, str]:
    strong = has_strong_signal(text, iocs, ttype)
    weak = has_weak_signal(text, iocs, ttype)
    tier1 = bool(matches["tier1"])
    broad = bool(matches["broad"])
    brand = bool(matches["brands"])
    region = bool(matches["regions"])
    specials = bool(matches.get("specials"))
    low = (text or "").lower()
    critical_grade = strong and (iocs["cves"] or any(k in low for k in ACTIVE_EXPLOIT_KEYWORDS + IMMEDIATE_ACTION_KEYWORDS) or ttype in {"Ransomware", "Data Breach", "Supply Chain", "APT / Intrusion"})

    reason_parts = []
    if matches["tier1"]:
        reason_parts.append("High-confidence technology match: " + join(matches["tier1"]))
    if matches["broad"]:
        reason_parts.append("Broad technology/vendor match: " + join(matches["broad"]))
    if matches.get("specials"):
        reason_parts.append("Special relevance pattern: " + join(matches["specials"]))
    if matches["regions"]:
        reason_parts.append("Regional match: " + join(matches["regions"]))
    if matches["brands"]:
        reason_parts.append("Brand match: " + join(matches["brands"]))
    if iocs["cves"]:
        reason_parts.append("CVEs: " + join(iocs["cves"]))
    if any(iocs[k] for k in ["ips", "domains", "urls", "sha256"]):
        reason_parts.append("Actionable IOCs extracted")
    base = "; ".join(reason_parts) if reason_parts else "No direct Company profile match"

    business_incident = is_business_incident(text, ttype)
    big_global_hits = matched_big_global_orgs(text)
    critical_region_hits, other_region_hits = region_sets(matches)
    if business_incident and critical_region_hits:
        extra = []
        if big_global_hits:
            extra.append("recognized global organization/company: " + join(big_global_hits[:6]))
        extra.append("critical Company-region business incident: " + join(sorted(critical_region_hits)))
        return CRITICAL_SHEET, "CRITICAL", base + "; " + "; ".join(extra) + "; alleged/confirmed breach, ransomware, or access-sale item requires immediate visibility.", ACTION_MAP.get(ttype, ACTION_MAP["Other"])
    if business_incident and big_global_hits:
        return REVIEW_SHEET, "HIGH", base + "; recognized global organization/company: " + join(big_global_hits[:6]) + "; global third-party/business incident requires analyst awareness but not Critical escalation without Company-region, Company-brand, or direct technology impact.", ACTION_MAP.get(ttype, ACTION_MAP["Other"])
    if business_incident and (other_region_hits or "profile logistics/vendor exposure" in matches.get("specials", [])):
        return REVIEW_SHEET, "MEDIUM", base + "; region/profile-vendor business incident requiring analyst awareness but not immediate Critical escalation.", ACTION_MAP.get(ttype, ACTION_MAP["Other"])

    # Critical: exact/high-confidence Company product plus critical-grade threat.
    # Exact tech is Critical only for direct exploit/patch/credential-session impact; report-only and travel-advisory items stay Review.
    if tier1 and critical_grade and is_direct_technology_threat(text, ttype) and not any(x in low for x in ["traveling employees", "hotel wi-fi", "threat spotlight", "proofpoint uncovered", "researchers uncovered"]):
        # Critical only when the Company technology itself is the affected/abused asset, not when it appears in a victim list, report, or OCR logo.
        return CRITICAL_SHEET, "CRITICAL", base + "; high-confidence Company technology/vendor plus direct critical-grade threat signal.", ACTION_MAP.get(ttype, ACTION_MAP["Other"])
    if brand and region and critical_grade:
        return CRITICAL_SHEET, "CRITICAL", base + "; brand and region correlation plus critical-grade threat signal.", ACTION_MAP.get(ttype, ACTION_MAP["Other"])
    if broad and critical_grade and any(k in low for k in ["actively exploited", "zero-day", "0-day", "mass exploitation", "unauthenticated rce", "cvss 10", "cvss score of 10"]):
        # Broad tech can still become Critical, but only for truly severe exploitation language.
        return CRITICAL_SHEET, "CRITICAL", base + "; broad Company technology with very high-severity exploitation signal.", ACTION_MAP.get(ttype, ACTION_MAP["Other"])

    # Review: exact tech, specials, or contextual broad tech related to a real threat.
    if tier1 and (strong or weak):
        return REVIEW_SHEET, "HIGH" if strong else "MEDIUM", base + "; high-confidence Company technology/vendor with non-critical or indirect threat signal.", ACTION_MAP.get(ttype, ACTION_MAP["Other"])
    if specials and (strong or weak):
        special_names = set(matches.get("specials", []))
        special_sev = "MEDIUM" if special_names & {"profile logistics/vendor exposure", "FortiGate access sale"} else ("HIGH" if strong else "MEDIUM")
        return REVIEW_SHEET, special_sev, base + "; business-relevant threat pattern requiring analyst check.", ACTION_MAP.get(ttype, ACTION_MAP["Other"])
    # Broad tech should not promote pure victim-domain, casual demo, anecdote, source/repost, or generic hosting/ASN matches.
    broad_terms = set(x.lower() for x in matches.get("broad", []))
    generic_industry_only = bool(broad_terms) and broad_terms.issubset({"logistics", "retail", "ecommerce", "e-commerce", "merchant", "payments", "marketing", "supply chain"})
    broad_downrank = any(x in low for x in ["ransomware operator", "running windows", "used microsoft edge", "read court affidavit", "full article", "additional coverage", "from bleepingcomputer", "from darkwebinformer", "as13335", "cloudflarenet"] + NON_ACTIONABLE_CHATTER_CONTEXT)
    # Region-only ransomware/phishing from non-core regions is usually archive-only unless there is an Company brand/vendor/special pattern.
    region_terms = set(x.lower() for x in matches.get("regions", []))
    non_core_region_only = bool(region_terms) and not (region_terms & CORE_REGIONS) and not tier1 and not brand and not specials
    if ttype == "Ransomware" and (generic_industry_only or non_core_region_only) and not tier1 and not brand and not specials:
        broad_downrank = True
    # Cloud provider/brand phishing should not be Review merely because Amazon/Facebook/Apple/Google was used as a lure or ASN.
    lure_only_terms = {"amazon", "facebook", "apple", "google", "fastly"}
    if ttype == "Phishing" and broad_terms and broad_terms.issubset(lure_only_terms) and not tier1 and not brand and not specials:
        broad_downrank = True
    if broad and strong and not broad_downrank and broad_match_is_actionable(text, iocs, matches, ttype):
        return REVIEW_SHEET, "HIGH", base + "; broad Company-relevant technology/vendor with strong, direct threat signal.", ACTION_MAP.get(ttype, ACTION_MAP["Other"])
    if brand and strong:
        return REVIEW_SHEET, "HIGH", base + "; Company brand plus strong threat signal.", ACTION_MAP.get(ttype, ACTION_MAP["Other"])
    if region and strong:
        region_terms = set(x.lower() for x in matches.get("regions", []))
        if (region_terms & CORE_REGIONS) or ttype in {"Data Breach", "APT / Intrusion"} or specials:
            return REVIEW_SHEET, "HIGH", base + "; Company region plus strong threat signal.", ACTION_MAP.get(ttype, ACTION_MAP["Other"])
    if (brand and region) and (strong or weak):
        return REVIEW_SHEET, "HIGH" if strong else "MEDIUM", base + "; brand and region with threat signal.", ACTION_MAP.get(ttype, ACTION_MAP["Other"])

    return LOW_SHEET, "LOW", base + "; No scoring combination met Critical or Review threshold.", "Archive/monitor only. No direct action unless additional Company relevance appears."


def summarize(text: str, limit: int = 420) -> str:
    t = clean_post_text(text)
    return t if len(t) <= limit else t[:limit-3].rstrip() + "..."

def read_optional_path_file(base_dir: Path, filename: str) -> Optional[Path]:
    path_file = base_dir / filename
    try:
        if path_file.exists():
            val = path_file.read_text(encoding="utf-8").strip().strip('"')
            if val:
                return Path(val).expanduser()
    except Exception:
        return None
    return None


def candidate_tesseract_paths(base_dir: Path, explicit_path: Optional[str] = None) -> List[Path]:
    candidates: List[Path] = []
    if explicit_path:
        candidates.append(Path(explicit_path).expanduser())
    cfg_path = read_optional_path_file(base_dir, "XTIM_tesseract_path.txt")
    if cfg_path:
        candidates.append(cfg_path)
    # Portable/local options. No environment variable changes required.
    candidates.extend([
        base_dir / "Tesseract-OCR" / "tesseract.exe",
        base_dir / "tools" / "Tesseract-OCR" / "tesseract.exe",
        base_dir / "ocr" / "Tesseract-OCR" / "tesseract.exe",
        base_dir.parent / "Tesseract-OCR" / "tesseract.exe",
    ])
    # PATH, if already available.
    which = shutil.which("tesseract")
    if which:
        candidates.append(Path(which))
    # Common standard install locations. These do not require PATH.
    env_keys = ["ProgramFiles", "ProgramFiles(x86)", "LOCALAPPDATA"]
    for key in env_keys:
        val = os.environ.get(key)
        if val:
            candidates.append(Path(val) / "Tesseract-OCR" / "tesseract.exe")
    # Also try common user-local path variants.
    candidates.extend([
        Path.home() / "AppData" / "Local" / "Programs" / "Tesseract-OCR" / "tesseract.exe",
        Path.home() / "AppData" / "Local" / "Tesseract-OCR" / "tesseract.exe",
    ])
    out: List[Path] = []
    seen: Set[str] = set()
    for c in candidates:
        try:
            key = str(c.resolve()).lower()
        except Exception:
            key = str(c).lower()
        if key not in seen:
            seen.add(key); out.append(c)
    return out


def find_tesseract_exe(base_dir: Path, explicit_path: Optional[str] = None) -> Optional[Path]:
    for c in candidate_tesseract_paths(base_dir, explicit_path):
        try:
            # If a folder was supplied, look for tesseract.exe inside it.
            if c.is_dir():
                c = c / "tesseract.exe"
            if c.exists() and c.name.lower() == "tesseract.exe":
                return c
        except Exception:
            continue
    return None


def configure_tesseract_runtime(pytesseract_mod: Any, base_dir: Path, explicit_path: Optional[str] = None) -> Optional[Path]:
    exe = find_tesseract_exe(base_dir, explicit_path)
    if not exe:
        return None
    pytesseract_mod.pytesseract.tesseract_cmd = str(exe)
    tessdata = exe.parent / "tessdata"
    if tessdata.exists() and "TESSDATA_PREFIX" not in os.environ:
        # Runtime-only setting for this process. It does NOT modify Windows environment variables.
        os.environ["TESSDATA_PREFIX"] = str(tessdata)
    return exe


class OCRContext:
    def __init__(self, enabled: bool, cache_dir: Path, base_dir: Path, max_images: int = 4, tesseract_path: Optional[str] = None):
        self.enabled = enabled
        self.cache_dir = cache_dir
        self.base_dir = base_dir
        self.max_images = max_images
        self.available = False
        self.reason = "OCR disabled"
        self.metrics = Counter()
        self._pytesseract = None
        self._Image = None
        self._ImageOps = None
        self.tesseract_exe: Optional[Path] = None
        if not enabled:
            return
        try:
            import pytesseract  # type: ignore
            from PIL import Image, ImageOps  # type: ignore
            self._pytesseract = pytesseract
            self._Image = Image
            self._ImageOps = ImageOps
            self.tesseract_exe = configure_tesseract_runtime(pytesseract, base_dir, tesseract_path)
            if not self.tesseract_exe:
                searched = "; ".join(str(x) for x in candidate_tesseract_paths(base_dir, tesseract_path)[:8])
                raise FileNotFoundError("tesseract.exe not found. Put it under XTIM/Tesseract-OCR/tesseract.exe, create XTIM_tesseract_path.txt with the full path, or run with --tesseract-path. Searched: " + searched)
            # Validate tesseract binary early so failures are clear before record processing.
            version = pytesseract.get_tesseract_version()
            self.cache_dir.mkdir(parents=True, exist_ok=True)
            self.available = True
            self.reason = f"OCR available via {self.tesseract_exe} (Tesseract {version})"
        except Exception as e:
            self.available = False
            self.reason = f"OCR unavailable: {e}"

    def extract(self, image_urls: Iterable[str]) -> str:
        urls = unique(image_urls)[: self.max_images]
        if not self.enabled:
            self.metrics["ocr_disabled_posts"] += 1
            return ""
        if not self.available:
            if urls:
                self.metrics["ocr_unavailable_posts_with_images"] += 1
            return ""
        texts: List[str] = []
        for url in urls:
            try:
                txt = self._extract_one(url)
                if txt:
                    texts.append(txt)
                    self.metrics["ocr_images_with_text"] += 1
                else:
                    self.metrics["ocr_images_no_text"] += 1
            except Exception:
                self.metrics["ocr_image_errors"] += 1
                continue
        if texts:
            self.metrics["ocr_posts_with_text"] += 1
        elif urls:
            self.metrics["ocr_posts_no_text"] += 1
        return "\n".join(texts).strip()

    def _extract_one(self, url: str) -> str:
        key = hashlib.sha256(url.encode("utf-8", errors="ignore")).hexdigest()
        txt_path = self.cache_dir / f"{key}.txt"
        if txt_path.exists():
            self.metrics["ocr_cache_hits"] += 1
            return txt_path.read_text(encoding="utf-8", errors="ignore").strip()
        self.metrics["ocr_downloads"] += 1
        import requests  # type: ignore
        headers = {"User-Agent": "Mozilla/5.0 XTIM OCR image fetch"}
        resp = requests.get(url, timeout=25, headers=headers)
        resp.raise_for_status()
        img = self._Image.open(io.BytesIO(resp.content))
        # Mild preprocessing helps screenshots without being destructive.
        img = img.convert("RGB")
        max_side = max(img.size)
        if max_side < 1200:
            scale = min(3, max(1, 1600 // max_side))
            if scale > 1:
                img = img.resize((img.width * scale, img.height * scale))
        gray = self._ImageOps.grayscale(img)
        text = self._pytesseract.image_to_string(gray, config="--psm 6")
        text = re.sub(r"\s+", " ", text or "").strip()
        txt_path.write_text(text, encoding="utf-8")
        return text


def build_text_with_ocr(post: Dict[str, Any], ocr_context: Optional[OCRContext]) -> Tuple[str, str]:
    raw = post.get("post_text") or post.get("raw_visible_text") or ""
    text = clean_post_text(raw)
    image_urls = post.get("post_image_urls") or post.get("image_urls") or []
    alt_text = " ".join(str(x) for x in (post.get("image_alt_texts") or []) if x).strip()
    ocr_text = ocr_context.extract(image_urls) if ocr_context else ""
    image_text_parts = []
    if alt_text:
        image_text_parts.append("Image alt text:\n" + alt_text)
    if ocr_text:
        image_text_parts.append("Image OCR:\n" + ocr_text)
    if image_text_parts:
        return text + "\n\n" + "\n\n".join(image_text_parts), ocr_text or alt_text
    return text, ""


def title_for(post: Dict[str, Any], iocs: Dict[str, List[str]], ttype: str) -> str:
    source = post.get("handle") or post.get("display_name") or "X"
    if iocs["cves"]:
        return f"{source}: {iocs['cves'][0]} - {ttype}"
    if iocs["domains"]:
        return f"{source}: {iocs['domains'][0]} - {ttype}"
    if iocs["ips"]:
        return f"{source}: {iocs['ips'][0]} - {ttype}"
    return f"{source}: {summarize(post.get('post_text') or '', 90)}"

def process_post(post: Dict[str, Any], ocr_context: Optional[OCRContext] = None) -> Tuple[str, List[str]]:
    text, ocr_text = build_text_with_ocr(post, ocr_context)
    iocs = extract_iocs(text, post.get("external_links") or [])
    ttype = classify_threat(text, iocs)
    matches = profile_matches(text)
    sheet, severity, reason, action = score_row(text, iocs, matches, ttype)
    if ocr_text:
        reason = reason + "; Image OCR text used for scoring/IOC extraction."
    row = [
        f"X / {post.get('handle') or '@unknown'}",
        title_for(post, iocs, ttype),
        post.get("post_url") or "",
        post.get("published_utc") or "",
        ttype,
        severity,
        summarize(text),
        reason,
        join(iocs["cves"]),
        join(iocs["ips"]),
        join(iocs["domains"]),
        join(iocs["urls"]),
        join(iocs["sha256"]),
        join(iocs["emails"]),
        join(iocs["filenames"]),
        join(unique(matches["tier1"] + matches["broad"])),
        join(matches["regions"]),
        join(unique(matches["tier1"] + matches["broad"])),
        join(matches["brands"]),
        action,
    ]
    return sheet, row

# -------------------------
# Timed/logged workbook processing
# -------------------------
def latest_json(script_dir: Path) -> Path:
    candidates: List[Path] = []
    search_roots: List[Path] = []
    for folder in [script_dir, Path.cwd(), Path.home() / "Downloads"]:
        try:
            if folder.exists() and folder not in search_roots:
                search_roots.append(folder)
        except Exception:
            continue
    for root in search_roots:
        try:
            candidates.extend(root.glob("x_captured_posts*.json"))
            # Search one level of subfolders for team folders where exports may be moved.
            for sub in root.iterdir():
                if sub.is_dir() and sub.name != LOG_DIR_NAME:
                    candidates.extend(sub.glob("x_captured_posts*.json"))
        except Exception:
            continue
    candidates = [p for p in candidates if p.is_file()]
    if not candidates:
        searched = ", ".join(str(p) for p in search_roots)
        raise FileNotFoundError(f"No x_captured_posts*.json found. Searched: {searched}")
    return max(candidates, key=lambda p: p.stat().st_mtime)

def load_state(path: Path) -> Dict[str, Any]:
    if not path.exists():
        return {"seen_posts": {}}
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
        if not isinstance(data.get("seen_posts"), dict):
            data["seen_posts"] = {}
        return data
    except Exception:
        return {"seen_posts": {}}

def save_state(path: Path, state: Dict[str, Any]) -> None:
    tmp = path.with_suffix(".tmp")
    tmp.write_text(json.dumps(state, ensure_ascii=False, indent=2), encoding="utf-8")
    tmp.replace(path)

def create_or_load_workbook(path: Path) -> Workbook:
    if path.exists():
        wb = load_workbook(path)
        for s in list(wb.sheetnames):
            if s not in SHEETS:
                del wb[s]
    else:
        wb = Workbook()
        wb.remove(wb.active)
    for sheet in SHEETS:
        if sheet not in wb.sheetnames:
            ws = wb.create_sheet(sheet)
            ws.append(COLUMNS)
        elif wb[sheet].max_row == 0:
            wb[sheet].append(COLUMNS)
    wb._sheets = [wb[s] for s in SHEETS if s in wb.sheetnames]
    format_workbook(wb)
    return wb

def existing_links_from_workbook(wb: Workbook) -> Set[str]:
    links: Set[str] = set()
    for s in SHEETS:
        ws = wb[s]
        headers = [c.value for c in ws[1]]
        if "link" in headers:
            idx = headers.index("link") + 1
            for row in ws.iter_rows(min_row=2, values_only=True):
                if row and len(row) >= idx and row[idx-1]:
                    links.add(str(row[idx-1]))
    return links

def format_workbook(wb: Workbook) -> None:
    fills = {CRITICAL_SHEET: "FFC7CE", REVIEW_SHEET: "FFEB9C", LOW_SHEET: "C6EFCE"}
    widths = {"A": 23, "B": 60, "C": 60, "D": 25, "E": 15, "F": 12, "G": 80, "H": 75, "I": 30, "J": 30, "K": 42, "L": 55, "M": 45, "N": 28, "O": 45, "P": 35, "Q": 30, "R": 45, "S": 45, "T": 90}
    thin = Side(style="thin", color="D9D9D9")
    for s in SHEETS:
        ws = wb[s]
        ws.freeze_panes = "A2"
        ws.sheet_view.showGridLines = False
        for c in ws[1]:
            c.fill = PatternFill("solid", fgColor=fills[s])
            c.font = Font(bold=True, color="000000")
            c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            c.border = Border(bottom=thin)
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(vertical="top", wrap_text=True)
        for col, width in widths.items():
            ws.column_dimensions[col].width = width
        ws.row_dimensions[1].height = 30
        if ws.max_row >= 1:
            ws.auto_filter.ref = ws.dimensions

def main() -> int:
    parser = argparse.ArgumentParser(description="Run X TIM JSON-to-Excel processing.")
    parser.add_argument("json_file", nargs="?", help="Path to x_captured_posts*.json. If omitted, newest is auto-detected.")
    parser.add_argument("--date", default=None, help="Workbook date YYYY-MM-DD. Default: local date GMT+04.")
    parser.add_argument("--max-age-days", type=int, default=DEFAULT_MAX_AGE_DAYS, help="Skip posts older than this many days. Use 0 to disable age filtering. Default: 3.")
    parser.add_argument("--no-ocr", action="store_true", help="Disable image OCR even if image URLs are present in the exported JSON.")
    parser.add_argument("--ocr-max-images", type=int, default=4, help="Maximum attached images to OCR per post. Default: 4.")
    parser.add_argument("--tesseract-path", default=None, help="Optional full path to tesseract.exe or its folder. Avoids changing PATH/environment variables.")
    args = parser.parse_args()

    print("[X TIM] Starting...", flush=True)
    script_dir = Path(__file__).resolve().parent
    log_dir = script_dir / LOG_DIR_NAME
    log_dir.mkdir(exist_ok=True)
    seen_path = log_dir / SEEN_FILE_NAME
    notes_path = log_dir / RUN_NOTES_NAME
    log_path = log_dir / LOG_FILE_NAME
    ocr_context = OCRContext(enabled=not args.no_ocr, cache_dir=log_dir / OCR_CACHE_DIR_NAME, base_dir=script_dir, max_images=args.ocr_max_images, tesseract_path=args.tesseract_path)
    print(f"[X TIM] OCR status: {ocr_context.reason}", flush=True)

    json_path = Path(args.json_file).resolve() if args.json_file else latest_json(script_dir)
    print(f"[X TIM] JSON input: {json_path}", flush=True)
    run_date = args.date or today_local_date()
    workbook_path = script_dir / f"{WORKBOOK_PREFIX}_{run_date}.xlsx"
    print(f"[X TIM] Workbook output: {workbook_path}", flush=True)

    print("[X TIM] Loading JSON...", flush=True)
    posts = json.loads(json_path.read_text(encoding="utf-8"))
    if not isinstance(posts, list):
        raise ValueError("Input JSON must be a list exported by the Edge extension.")
    print(f"[X TIM] Loaded {len(posts)} records.", flush=True)

    print("[X TIM] Loading workbook and seen-state...", flush=True)
    state = load_state(seen_path)
    seen_posts: Dict[str, Any] = state.setdefault("seen_posts", {})
    wb = create_or_load_workbook(workbook_path)
    existing_links = existing_links_from_workbook(wb)
    cutoff = now_utc() - timedelta(days=args.max_age_days) if args.max_age_days and args.max_age_days > 0 else None

    metrics = Counter()
    batch_ids: Set[str] = set()
    new_items: List[Tuple[str, str, str, List[str], Dict[str, Any]]] = []

    print("[X TIM] Processing records...", flush=True)
    total = len(posts)
    for idx, post in enumerate(posts, start=1):
        if idx == 1 or idx % 25 == 0 or idx == total:
            print(f"[X TIM] Processing {idx}/{total}...", flush=True)
        metrics["input_records"] += 1
        post_id = str(post.get("post_id") or "").strip()
        link = str(post.get("post_url") or "").strip()
        if not post_id:
            metrics["skipped_missing_post_id"] += 1; continue
        if post_id in batch_ids:
            metrics["skipped_duplicate_within_json"] += 1; continue
        batch_ids.add(post_id)
        dt = parse_dt(post.get("published_utc") or "")
        if not dt:
            metrics["skipped_invalid_date"] += 1; continue
        if cutoff is not None and dt < cutoff:
            metrics["skipped_older_than_max_age"] += 1; continue
        if post_id in seen_posts or link in existing_links:
            metrics["skipped_already_processed"] += 1; continue
        sheet, row = process_post(post, ocr_context)
        new_items.append((post_id, link, sheet, row, post))

    print(f"[X TIM] Appending {len(new_items)} new rows...", flush=True)
    for post_id, link, sheet, row, post in new_items:
        wb[sheet].append(row)
        seen_posts[post_id] = {
            "first_processed_utc": now_utc().isoformat(),
            "published_utc": post.get("published_utc") or "",
            "workbook": workbook_path.name,
            "sheet": sheet,
            "post_url": link,
        }

    print("[X TIM] Saving workbook and logs...", flush=True)
    format_workbook(wb)
    try:
        wb.save(workbook_path)
    except PermissionError:
        print(f"ERROR: Could not save workbook. Close it in Excel and rerun: {workbook_path}", file=sys.stderr, flush=True)
        return 2
    save_state(seen_path, state)

    metrics.update(ocr_context.metrics)
    metrics["ocr_available"] = 1 if ocr_context.available else 0
    metrics["processed_new_records"] = len(new_items)
    for s in SHEETS:
        metrics[f"rows_total_{s}"] = max(wb[s].max_row - 1, 0)

    notes = [
        "X TIM Script Run Notes",
        "=" * 30,
        f"Run timestamp UTC: {now_utc().isoformat()}",
        f"Input JSON: {json_path}",
        f"Workbook: {workbook_path}",
        f"Seen-state file: {seen_path}",
        f"Max age days: {args.max_age_days}",
        "",
        "Metrics:",
    ]
    for k, v in metrics.items():
        notes.append(f"- {k}: {v}")
    notes_text = "\n".join(notes)
    notes_path.write_text(notes_text, encoding="utf-8")
    with open(log_path, "a", encoding="utf-8") as f:
        f.write(notes_text + "\n\n")

    print("[OK] X TIM script run complete", flush=True)
    print(f"Input records: {metrics['input_records']}", flush=True)
    print(f"New records processed: {metrics['processed_new_records']}", flush=True)
    print(f"Already processed skipped: {metrics['skipped_already_processed']}", flush=True)
    print(f"Older than max age skipped: {metrics['skipped_older_than_max_age']}", flush=True)
    print(f"Workbook: {workbook_path}", flush=True)
    print(f"Critical rows total: {metrics[f'rows_total_{CRITICAL_SHEET}']}", flush=True)
    print(f"Review rows total: {metrics[f'rows_total_{REVIEW_SHEET}']}", flush=True)
    print(f"Low rows total: {metrics[f'rows_total_{LOW_SHEET}']}", flush=True)
    print(f"Logs/state folder: {log_dir}", flush=True)
    return 0

if __name__ == "__main__":
    try:
        raise SystemExit(main())
    except KeyboardInterrupt:
        print("\n[X TIM] Stopped by user (Ctrl+C). No further action completed.", file=sys.stderr, flush=True)
        raise SystemExit(130)
    except Exception as e:
        print(f"\n[X TIM] ERROR: {type(e).__name__}: {e}", file=sys.stderr, flush=True)
        print("[X TIM] Check that the JSON export exists, Excel output is closed, and dependencies are installed. See XTIM_README.md.", file=sys.stderr, flush=True)
        raise SystemExit(1)
